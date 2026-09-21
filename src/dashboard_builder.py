"""
Gera o dapl_dashboard.html a partir do Excel consolidado, reaproveitando o
template HTML/JS já validado manualmente (templates/dashboard_template.html
e templates/dashboard.js). Este módulo só cuida de extrair os dados do
Excel e injetá-los — a lógica visual/interativa do dashboard não muda.
"""
import json
import os
from openpyxl import load_workbook

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "..", "templates")


def _sheet_to_records(ws):
    headers = [c.value for c in ws[1]]
    key_col = headers.index("Chave da item")
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[key_col] is None:
            continue
        records.append({h: v for h, v in zip(headers, row)})
    return records


def _sheet_to_matrix(ws):
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
        rows.append(list(row))
    return rows


def build_dashboard(excel_path: str, output_path: str):
    wb = load_workbook(excel_path, data_only=True)
    records = _sheet_to_records(wb["Indicadores DAPL"])
    kpi_matrix = _sheet_to_matrix(wb["KPIs Mensais"])

    payload = {"records": records, "kpi": kpi_matrix}
    payload_json = json.dumps(payload, ensure_ascii=False, default=str)

    with open(os.path.join(TEMPLATE_DIR, "dashboard_template.html"), encoding="utf-8") as f:
        template = f.read()
    with open(os.path.join(TEMPLATE_DIR, "dashboard.js"), encoding="utf-8") as f:
        script = f.read()

    html = template.replace("__PAYLOAD_JSON__", payload_json).replace("__DASHBOARD_JS__", script)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)

    return output_path
