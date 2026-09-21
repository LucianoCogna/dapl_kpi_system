"""
Gera o arquivo Indicadores_DAPL_Consolidado.xlsx a partir da lista de
métricas por card (saída de kpi_engine.compute_card_metrics).

Reproduz a estrutura de duas abas usada durante o processamento manual:
- "Indicadores DAPL": uma linha por card
- "KPIs Mensais": indicador x mês, com fórmulas (não valores fixos) para que
  o arquivo continue recalculando se alguém editar uma célula depois.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import config

HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
FLAG_FONT = Font(name="Arial", size=10, bold=True, color="1F4E78")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    "Tipo de item", "Chave da item", "Resumo", "Status", "Responsável",
    "Mês de Referência", "Mês Fim Refinamento",
    "Flag Engenharia (DENA)", "Flag Plataforma (DDPL)", "Flag DataViz (DCOD)", "Flag Governança (DGOD)",
    "Flag Novo Produto de Dados", "Flag Evolução",
    "Tempo em Refinamento (dias)", "Validar Tempo Refinamento",
    "Time to Market (dias corridos)", "Tempo em Homologação (dias)",
    "Classificação Tamanho", "Elegível TTM Previsto?", "Tempo Bloqueio Externo (dias)",
    "Qtde Dias Líquidos", "Limite (dias)", "Dentro do Prazo?",
    "Changelog incompleto?",
]


def _col_letter(name):
    return get_column_letter(COLUMNS.index(name) + 1)


def build_workbook(records: list[dict], months: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Indicadores DAPL"

    # Header
    for i, name in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    # Rows
    for r, rec in enumerate(records, start=2):
        for i, col_name in enumerate(COLUMNS, 1):
            val = rec.get(col_name)
            c = ws.cell(row=r, column=i, value=val)
            c.font = FLAG_FONT if col_name.startswith("Flag") else BODY_FONT
            if col_name.startswith("Flag") or col_name in ("Mês de Referência", "Mês Fim Refinamento"):
                c.alignment = Alignment(horizontal="center")
            c.border = BORDER

    last_row = len(records) + 1
    ws.freeze_panes = "A2"
    for i in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    # ---------------- KPIs Mensais ----------------
    kpi = wb.create_sheet("KPIs Mensais")
    month_labels = [m.capitalize() for m in months]
    kpi.cell(row=1, column=1, value="Indicador").font = HEADER_FONT
    kpi.cell(row=1, column=1).fill = HEADER_FILL
    kpi.cell(row=1, column=2, value="Periodicidade").font = HEADER_FONT
    kpi.cell(row=1, column=2).fill = HEADER_FILL
    for i, label in enumerate(month_labels, 3):
        c = kpi.cell(row=1, column=i, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL

    def month_col_letters():
        return [get_column_letter(i) for i in range(3, 3 + len(month_labels))]

    mref = f"'Indicadores DAPL'!${_col_letter('Mês de Referência')}$2:${_col_letter('Mês de Referência')}${last_row}"
    mfim = f"'Indicadores DAPL'!${_col_letter('Mês Fim Refinamento')}$2:${_col_letter('Mês Fim Refinamento')}${last_row}"

    def flag_range(col_name):
        return f"'Indicadores DAPL'!${_col_letter(col_name)}$2:${_col_letter(col_name)}${last_row}"

    kpi_rows = [
        ("Total de Demandas", lambda col: f"=COUNTIF({mref},UPPER({col}$1))"),
        ("% Novo Produto", lambda col: f"=IF({col}$2=0,0,COUNTIFS({mref},UPPER({col}$1),{flag_range('Flag Novo Produto de Dados')},1)/{col}$2)"),
        ("% Evolução Produto", lambda col: f"=IF({col}$2=0,0,COUNTIFS({mref},UPPER({col}$1),{flag_range('Flag Evolução')},1)/{col}$2)"),
        ("% Engenharia de Dados", lambda col: f"=IF({col}$2=0,0,COUNTIFS({mref},UPPER({col}$1),{flag_range('Flag Engenharia (DENA)')},1)/{col}$2)"),
        ("% Plataforma", lambda col: f"=IF({col}$2=0,0,COUNTIFS({mref},UPPER({col}$1),{flag_range('Flag Plataforma (DDPL)')},1)/{col}$2)"),
        ("% DataViz", lambda col: f"=IF({col}$2=0,0,COUNTIFS({mref},UPPER({col}$1),{flag_range('Flag DataViz (DCOD)')},1)/{col}$2)"),
        ("% Governança", lambda col: f"=IF({col}$2=0,0,COUNTIFS({mref},UPPER({col}$1),{flag_range('Flag Governança (DGOD)')},1)/{col}$2)"),
        ("Tempo para insight (méd. dias em Refinamento) — apenas Novo Produto de Dados",
         lambda col: f"=IFERROR(AVERAGEIFS({flag_range('Tempo em Refinamento (dias)')},{mfim},UPPER({col}$1),{flag_range('Flag Novo Produto de Dados')},1),\"Sem dados\")"),
        ("Time to Market (dias corridos, Refinamento até Aceito) — apenas Novo Produto de Dados",
         lambda col: f"=IFERROR(AVERAGEIFS({flag_range('Time to Market (dias corridos)')},{mref},UPPER({col}$1),{flag_range('Flag Novo Produto de Dados')},1),\"Sem dados\")"),
        ("Tempo em Homologação (dias, Concluído até Aceito)",
         lambda col: f"=IFERROR(AVERAGEIFS({flag_range('Tempo em Homologação (dias)')},{mref},UPPER({col}$1)),\"Sem dados\")"),
        ("% Novos Produtos entregues dentro do TTM previsto (por classificação P/M/G)",
         lambda col: (f"=IFERROR(COUNTIFS({mref},UPPER({col}$1),{flag_range('Elegível TTM Previsto?')},\"Sim\","
                       f"{flag_range('Dentro do Prazo?')},\"Sim\")"
                       f"/COUNTIFS({mref},UPPER({col}$1),{flag_range('Elegível TTM Previsto?')},\"Sim\"),\"Sem dados\")")),
    ]

    for r, (label, formula_fn) in enumerate(kpi_rows, start=2):
        kpi.cell(row=r, column=1, value=label).font = BODY_FONT
        kpi.cell(row=r, column=2, value="Mensal").font = BODY_FONT
        for col in month_col_letters():
            kpi.cell(row=r, column=column_index_from(col), value=formula_fn(col)).font = BODY_FONT

    kpi.column_dimensions["A"].width = 55
    for col in month_col_letters():
        kpi.column_dimensions[col].width = 14

    return wb


def column_index_from(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)
